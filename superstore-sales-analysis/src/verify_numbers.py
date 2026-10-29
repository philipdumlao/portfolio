
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
df = pd.read_excel(ROOT / "data" / "sample_-_superstore.xls", sheet_name="Orders")
wb = load_workbook(ROOT / "excel" / "Superstore_Sales_Performance.xlsx", data_only=True)
bad = 0

def chk(label, got, exp, tol=1e-6):
    global bad
    ok = abs(got - exp) <= tol * max(1, abs(exp))
    bad += not ok
    print(("OK  " if ok else "FAIL"), label, got, exp)

k = wb["KPI_Summary"]
chk("Total Sales", k["B5"].value, df.Sales.sum())
chk("Total Profit", k["B6"].value, df.Profit.sum())
chk("Margin", k["B7"].value, df.Profit.sum() / df.Sales.sum())
chk("Quantity", k["B8"].value, df.Quantity.sum())
chk("Avg Discount", k["B9"].value, df.Discount.mean())
chk("Loss lines", k["B11"].value, (df.Profit < 0).sum())
chk("Correlation", k["B13"].value, df.Discount.corr(df.Profit))

def table(ws, first, last, dim):
    g = df.groupby(dim).agg(S=("Sales", "sum"), P=("Profit", "sum"), Q=("Quantity", "sum"), D=("Discount", "mean"))
    for r in range(first, last + 1):
        n = ws.cell(r, 1).value
        chk(f"{ws.title} {n} sales", ws.cell(r, 2).value, g.S[n])
        chk(f"{ws.title} {n} profit", ws.cell(r, 3).value, g.P[n])
        chk(f"{ws.title} {n} qty", ws.cell(r, 5).value, g.Q[n])
        chk(f"{ws.title} {n} avgdisc", ws.cell(r, 6).value, g.D[n])

table(wb["Pivot_Region"], 5, 8, "Region")
table(wb["Pivot_Category"], 5, 7, "Category")
table(wb["Pivot_SubCategory"], 5, 21, "Sub-Category")
pp = wb["Pivot_Product"]
g = df.groupby("Product Name").agg(S=("Sales", "sum"), P=("Profit", "sum"))
for r in list(range(6, 16)) + list(range(19, 29)):
    n = pp.cell(r, 1).value
    chk(f"Product {n[:30]} profit", pp.cell(r, 4).value, g.P[n])
yr = df.groupby(df["Order Date"].dt.year).Profit.sum()
for r in range(5, 9):
    chk(f"Year {wb['Pivot_Year'].cell(r,1).value}", wb["Pivot_Year"].cell(r, 3).value, yr[wb["Pivot_Year"].cell(r, 1).value])
d = wb["Discount_Analysis"]
labs = pd.cut(df.Discount, [-1, 0, .1, .2, .3, .5, 1], labels=["No discount", "Up to 10%", "Up to 20%", "Up to 30%", "Up to 50%", "Over 50%"])
gb = df.groupby(labs, observed=True).agg(n=("Sales", "size"), P=("Profit", "sum"))
for r in range(5, 11):
    n = d.cell(r, 1).value
    chk(f"Band {n} lines", d.cell(r, 2).value, gb.n[n]); chk(f"Band {n} profit", d.cell(r, 4).value, gb.P[n])
print("\nFAILURES:", bad)
